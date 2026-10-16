"""Ekran görüntülerindeki formatta örnek Excel dosyaları üretir (Eylül 2026).

Kullanım:  python ornek_veri.py [klasör]      (varsayılan: ornek/)
Değerler ekran görüntülerinden alınmıştır, sadece denemek içindir.
"""
import sys
from pathlib import Path

from openpyxl import Workbook

MASTER = [
    # Dönem, Yıl, Sözleşme, KDV, Poz, Miktar, KDV'siz, Birim fiyat, KDV, $
    ("Ağustos", 2026, "MADEN", "Y", "Delme-Patlatma-Kazı-Yükleme-3.5 km'ye kadar Taşıma -Depolama", 900000, 2619000, 2.91, 314280, 2933280),
    ("Eylül", 2026, "MADEN", "Y", "Delme-Patlatma-Kazı-Yükleme-3.5 km'ye kadar Taşıma -Depolama", 1010248, 2939821.68, 2.91, 352778.60, 3292600.28),
    ("Eylül", 2026, "MADEN", "Y", "Delme-Kazı-Yükleme-3.5 km'ye kadar Taşıma -Depolama", 297489, 716948.49, 2.41, 86033.82, 802982.31),
    ("Eylül", 2026, "MADEN", "Y", "Yükleme-1 km'ye kadar Taşıma-Depolama", 8100, 10530, 1.30, 1263.60, 11793.60),
    ("Eylül", 2026, "MADEN", "Y", "Kiralık", None, 91199.68, 0, 10943.96, 102143.64),
    ("Eylül", 2026, "MADEN", "N", "Diğer", None, 99919.52, 0, 0, 99919.52),
    ("Eylül", 2026, "MADEN", "N", "Elektrik", None, -13231.89, 0, 0, -13231.89),
    ("Eylül", 2026, "MADEN", "N", "Yemek", None, 10744.20, 0, 0, 10744.20),
    ("Eylül", 2026, "MADEN", "N", "Prim", None, 0, 0, 0, 0),
    ("Eylül", 2026, "MADEN", "N", "Yakıt Farkı", None, 461614.65, 0, 0, 461614.65),
    ("Eylül", 2026, "MADEN", "Y", "Delme-Kazı-Yükleme-3.5 km'ye kadar Taşıma -Depolama Ocak Dışı", 45599, 109893.59, 2.41, 13187.23, 123080.82),
    ("Eylül", 2026, "MADEN", "Y", "Delme-Patlatma-Kazı-Yükleme-3.5 km'ye kadar Taşıma -Depolama Ocak Dışı", 1048, 3049.68, 2.91, 365.96, 3415.64),
    ("Eylül", 2026, "TSF", "Y", "Patlatmasız Kazı", 3275, 6877.50, 2.10, 825.30, 7702.80),
    ("Eylül", 2026, "TSF", "Y", "TSF Faz 2 Dolgu", 250199, 80063.68, 0.32, 9607.64, 89671.32),
    ("Eylül", 2026, "TSF", "Y", "Nebati metrekare", None, 0, 0, 0, 0),
    ("Eylül", 2026, "TSF", "Y", "Kiralık", None, 146655.34, 0, 17598.64, 164253.98),
    ("Eylül", 2026, "TSF", "N", "Diğer", None, 0, 0, 0, 0),
    ("Eylül", 2026, "RC", "Y", "RC_1 HANFA", None, 0, 0, 0, 0),
    ("Eylül", 2026, "RC", "Y", "RC-3 EPİROC D65", 14288, 228608, 16, 27432.96, 256040.96),
    ("Eylül", 2026, "RC", "Y", "Kiralık", None, 3512, 0, 421.44, 3933.44),
]

EKIPMAN = [
    # kod, ad, tip, company, birincil
    ("10-303", "ZE500G-3", "ESKAVATOR", "Pitworks", "Evet"),
    ("10-401", "PC850-1", "ESKAVATOR", "Pitworks", "Evet"),
    ("13-101", "D85-1", "DOZER", "Pitworks", "Evet"),
    ("15-504", "RC-4", "DELICI", "Pitworks", "Evet"),
    ("15-503", "RC-3", "RC", "Pitworks", "Evet"),
    ("15-301", "D55", "DELICI", "Pitworks", "Evet"),
    ("11-002", "DT-02", "KAMYON", "Pitworks", "Evet"),
    ("11-004", "DT-04", "KAMYON", "Pitworks", "Evet"),
    ("11-011", "DT-11", "KAMYON", "Pitworks", "Evet"),
    ("19-001", "D-MAX-1", "PIKAP", "Pitworks", "Hayır"),
    ("01 027 HA", "SHANTU DOZER", "DOZER", "Rental", "Hayır"),
    ("01 961 SNA", "HOWO 01 961 SNA", "HOWO KAMYON", "Tomdi", "Hayır"),
]

MAZOT = [
    # gün, kod, toplam litre
    (5, "10-303", 497), (5, "10-401", 963), (5, "13-101", 388), (5, "15-504", 1687), (5, "15-503", 1629),
    (5, "15-301", 983), (5, "11-002", 218), (5, "11-004", 234), (5, "11-011", 454), (5, "19-001", 69),
    (5, "01 027 HA", 76), (5, "01 961 SNA", 200),
    (6, "10-303", 593), (6, "11-002", 247), (6, "11-004", 240), (6, "11-011", 229), (6, "01 961 SNA", 300),
]

DEPO = [
    # tarih, kod, ru, tr, birim, miktar, fiyat, tutar, garaj, ekipman adı, grup, alt grup, gider yeri
    ("2026-09-01", 241, "Насос для минеральная масла", "YAĞ POMPASI", "piece", 2, 283050, 566100, "ATÖLYE", "ATÖLYE", "DİĞER STOKLAR", "ATÖLYE DİĞER STOK MALZEMELERİ", "ATÖLYE SARF MALZEME GİDERLERİ"),
    ("2026-09-01", 237, "Строп текстильный", "BEZ SAPAN 10 TON", "piece", 4, 484685.71, 1938742.86, "ATÖLYE", "ATÖLYE", "DİĞER STOKLAR", "ATÖLYE DİĞER STOK MALZEMELERİ", "ATÖLYE SARF MALZEME GİDERLERİ"),
    ("2026-09-01", 309, "Аккумулятор 6СТ 74Ah", "AKÜ 74AH VARIAN", "piece", 1, 892000, 892000, "19-001", "ISUZU D-MAX", "YEDEK PARÇA", "BİNEK ARAÇ YEDEK PARÇA", "BİNEK ARAÇ YEDEK PARÇA GİDERLERİ"),
    ("2026-09-01", 320, "сигнальная лента", "25MT UYARI LEVHASI", "piece", 1, None, 0, "İSG", "İSG", "DİĞER STOKLAR", "İSG DİĞER STOK MALZEMELERİ", "İSG SARF MALZEME GİDERLERİ"),
    ("2026-09-01", 319, "Светодиодный флаг", "LED GÜVENLİK BAYRAĞI", "piece", 1, 58035.71, 58035.71, "11-002", "MERCEDES 4851", "DİĞER STOKLAR", "KAMYON DİĞER STOK", "KAMYON SARF MALZEME GİDERLERİ"),
    ("2026-09-01", 319, "Светодиодный флаг", "LED GÜVENLİK BAYRAĞI", "piece", 1, 58035.71, 58035.71, "11-004", "MERCEDES 4851", "DİĞER STOKLAR", "KAMYON DİĞER STOK", "KAMYON SARF MALZEME GİDERLERİ"),
    ("2026-09-02", 89, "ХОМУТ 44-64", "METAL KELEPÇE 44-64", "piece", 9, 6400, 57600, "15-301", "FSL-600", "DİĞER STOKLAR", "İŞ MAKİNESİ DİĞER STOK", "DELİCİ SARF MALZEME GİDERLERİ"),
    ("2026-09-02", 90, "ХОМУТ 32-44", "METAL KELEPÇE 32-44", "piece", 18, 5500, 99000, "10-303", "ZE500G-3", "DİĞER STOKLAR", "İŞ MAKİNESİ DİĞER STOK", "EKSKAVATÖR SARF MALZEME GİDERLERİ"),
]

PERSONEL = [
    ("Burg'ulash qurilmasi mexanigi/Baş Operatör ( ROC)", "Delme ve Patlatma", 1, 11865),
    ("Bosh Yo'riqchi/Baş Şoför", "Operasyon", 2, 18144),
    ("Ish boshqaruvchi/Saha Formeni", "Operasyon", 4, 16217),
    ("Kon muhandisi/Maden mühendisi", "Operasyon", 3, 15616),
    ("Loyiha Menejeri/Proje müdürü", "Yönetim", 1, 22590),
    ("Ombor mudiri /Depo sorumlusu", "Atölye", 1, 18469),
    ("Portlatma muhandisi / Patlatma mühendisi", "Delme ve Patlatma", 2, 28658),
    ("Direktor / Müdür", "Yönetim", 1, 14961),
    ("Burg'ilash qurilmasi mashinisti / Delici Operatörü", "Delme ve Patlatma", 12, 10260),
    ("Samosval haydovchisi / Kamyon şoför", "Operasyon", 136, 14237),
    ("Ekskavator mashinisti / Ekskavatör operatörü", "Operasyon", 36, 12112),
    ("Kon muhandisi/Maden mühendisi", "Operasyon", 3, 29928),
    ("Moliyachi / Finansçı", "Idari Isler", 1, 16211),
]


def kaydet(wb, yol):
    wb.save(yol)
    print("yazıldı:", yol)


def uret(klasor: Path) -> dict:
    klasor.mkdir(parents=True, exist_ok=True)
    yollar = {}

    wb = Workbook(); ws = wb.active; ws.title = "Master Data"
    ws.append(["Dönem", "Yıl", "Hakediş Sözleşmesi", "KDV Y/N", "Hakediş Poz", "Yapılan iş m³ - m", "KDV'siz", "Birim Fiyat $", "KDV", "$"])
    for r in MASTER:
        ws.append(list(r))
    yollar["master"] = klasor / "Master_Data.xlsx"; kaydet(wb, yollar["master"])

    wb = Workbook(); ws = wb.active; ws.title = "Ekipman"
    ws.append(["Company No", "Equipment No", "Equipment Type", "Company", "Birincil Ekipman"])
    for r in EKIPMAN:
        ws.append(list(r))
    yollar["ekipman"] = klasor / "Ekipman_Listesi.xlsx"; kaydet(wb, yollar["ekipman"])

    bilgi = {k: (a, t, c) for k, a, t, c, _ in EKIPMAN}
    wb = Workbook(); ws = wb.active; ws.title = "Fuel"
    ws.append(["Week", "Day", "Month", "Year", "Company", "Company No", "Equipment No", "Equipment Type", "Not", "Fuel", "Fuel", "Toplam"])
    for gun, kod, litre in MAZOT:
        ad, tip, comp = bilgi[kod]
        ws.append([36, gun, 9, 2026, comp, kod, ad, tip, None, litre, None, litre])
    ws.append([40, 1, 10, 2026, "Pitworks", "11-002", "DT-02", "KAMYON", None, 300, None, 300])  # başka ay
    yollar["mazot"] = klasor / "Mazot.xlsx"; kaydet(wb, yollar["mazot"])

    wb = Workbook(); ws = wb.active; ws.title = "Depo"
    ws.append(["Дата", "Код", "Наименование Материала", None, "Ед.изм.", "кол.во", "Цена", "Сумма", "Гараж номер а/с",
               "Ekipman Adı", "Grup", "Alt Grup", "Gider Yeri"])
    for r in DEPO:
        ws.append(list(r))
    yollar["depo"] = klasor / "Depo.xlsx"; kaydet(wb, yollar["depo"])

    wb = Workbook(); ws = wb.active; ws.title = "Personel"
    ws.cell(1, 16, "Görev"); ws.cell(1, 17, "Birim"); ws.cell(1, 18, "Adet"); ws.cell(1, 19, "Birim Maaş"); ws.cell(1, 20, "Toplam Maaş")
    for c in range(16, 21):
        ws.cell(2, c, 0)
    for i, (g, b, a, m) in enumerate(PERSONEL, start=3):
        for c, v in zip(range(16, 21), (g, b, a, m, a * m)):
            ws.cell(i, c, v)
    yollar["personel"] = klasor / "Personel.xlsx"; kaydet(wb, yollar["personel"])
    return yollar


if __name__ == "__main__":
    uret(Path(sys.argv[1] if len(sys.argv) > 1 else "ornek"))
